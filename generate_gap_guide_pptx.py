#!/usr/bin/env python3
"""16:9 PowerPoint of GAP-GUIDE-2026-001 from the HTML guideline."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
LOGO = ROOT.parent / "gasiapac-official-doc-assets" / "g-asiapacific-logo-black.png"
OUT = ROOT / "GAP-GUIDE-2026-001-APN-Opportunity-Lark-Base.pptx"

NAVY = RGBColor(0x1B, 0x2A, 0x4A)
CREAM = RGBColor(0xF4, 0xEF, 0xE4)
PAPER = RGBColor(0xFF, 0xFD, 0xF8)
ORANGE = RGBColor(0xE8, 0x77, 0x22)
MUTE = RGBColor(0x5A, 0x56, 0x4E)
INK = RGBColor(0x1C, 0x24, 0x30)
WHITE = RGBColor(0xFF, 0xFB, 0xF7)
OK = RGBColor(0x1F, 0x7A, 0x3A)
BAD = RGBColor(0xB4, 0x23, 0x18)
OK_BG = RGBColor(0xDF, 0xF3, 0xE3)
BAD_BG = RGBColor(0xFF, 0xE1, 0xDC)
WARN_BG = RGBColor(0xFF, 0xE8, 0xD4)

W, H = Inches(13.333), Inches(7.5)


def rgb_hex(c: RGBColor) -> str:
    return f"{c[0]:02X}{c[1]:02X}{c[2]:02X}"


def fill(shape, color: RGBColor):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def set_run(run, text, size=14, bold=False, color=INK, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_text(slide, l, t, w, h, text, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT, font="Calibri", anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set(qn("a:anchor"), {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    set_run(run, text, size, bold, color, font)
    return box


def add_para_box(slide, l, t, w, h, lines, size=13, color=INK, bold_first=False, anchor=MSO_ANCHOR.TOP):
    """lines: list of str, or (text, bold, size, color) tuples."""
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(6)
        if isinstance(line, tuple):
            text, is_bold, sz, col = line
        else:
            text, is_bold, sz, col = line, (bold_first and i == 0), size, color
        run = p.add_run()
        set_run(run, text, sz, is_bold, col)
    return box


def card(slide, l, t, w, h, fill_c=PAPER):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    fill(sh, fill_c)
    sh.adjustments[0] = 0.08
    return sh


def kicker(slide, l, t, w, text, color=ORANGE):
    add_text(slide, l, t, w, Inches(0.28), text.upper(), size=11, bold=True, color=color, font="Calibri")


def footer(slide, light=False):
    col = RGBColor(0xC8, 0xC0, 0xB4) if light else RGBColor(0x8A, 0x84, 0x78)
    add_text(
        slide,
        Inches(0.5),
        Inches(7.18),
        Inches(12.3),
        Inches(0.24),
        "Internal  ·  v1.1  ·  © G-AsiaPacific",
        size=10,
        color=col,
    )


def cream_bg(slide):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    fill(sh, CREAM)
    sp = sh._element
    sp.getparent().remove(sp)
    slide.shapes._spTree.insert(2, sp)


def navy_bg(slide):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    fill(sh, NAVY)
    sp = sh._element
    sp.getparent().remove(sp)
    slide.shapes._spTree.insert(2, sp)


def pic(slide, path, l, t, w, h):
    if Path(path).exists():
        slide.shapes.add_picture(str(path), l, t, w, h)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]

    # 1 Cover
    s = prs.slides.add_slide(blank)
    navy_bg(s)
    pic(s, ASSETS / "apn-kt-cover.png", Inches(6.6), 0, Inches(6.74), H)
    overlay = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(7.4), H)
    fill(overlay, NAVY)
    if LOGO.exists():
        s.shapes.add_picture(str(LOGO), Inches(0.55), Inches(0.4), Inches(1.7), Inches(0.48))
    kicker(s, Inches(0.55), Inches(1.15), Inches(6), "Official documentation  ·  v1.1  ·  Issued")
    add_text(s, Inches(0.55), Inches(1.6), Inches(6.4), Inches(2.2), "APN Opportunity — Lark Base guidelines", size=36, bold=True, color=WHITE)
    add_text(
        s,
        Inches(0.55),
        Inches(4.0),
        Inches(6.2),
        Inches(1.4),
        "From 18 September 2026, type the opportunity in Lark. Lark still sends it to APN. Do not create a live customer opportunity until then.",
        size=16,
        color=RGBColor(0xE8, 0xDF, 0xD2),
    )
    add_text(s, Inches(0.55), Inches(5.7), Inches(6.2), Inches(0.8), "GAP sales  ·  Malaysia, Indonesia, Vietnam\nKnowledge Transfer 14 September 2026", size=13, color=ORANGE)
    footer(s, light=True)

    # 2 Why
    s = prs.slides.add_slide(blank)
    cream_bg(s)
    pic(s, ASSETS / "apn-kt-why.png", 0, 0, Inches(5.4), H)
    kicker(s, Inches(5.8), Inches(0.7), Inches(6.8), "Why Lark")
    add_text(s, Inches(5.8), Inches(1.15), Inches(6.8), Inches(1.2), "A different door into the same APN", size=28, bold=True, color=NAVY)
    add_para_box(
        s,
        Inches(5.8),
        Inches(2.5),
        Inches(6.8),
        Inches(3.6),
        [
            ("AWS launched a new APN. That new APN has limits for our team.", False, 16, INK),
            ("You type the opportunity in Lark. Lark still sends it to APN.", False, 16, INK),
            ("You are not keeping a second set of books.", True, 16, NAVY),
        ],
    )
    footer(s)

    # 3 Tables
    s = prs.slides.add_slide(blank)
    cream_bg(s)
    kicker(s, Inches(0.5), Inches(0.32), Inches(12), "What you will see")
    add_text(s, Inches(0.5), Inches(0.62), Inches(12), Inches(0.5), "Only Submission is typed. Ignore Client management and funding.", size=22, bold=True, color=NAVY)
    tables = [
        ("APN Oppt Submission Table", "The only table you type into. Draft with Grid + Add Record. Send from Form. Watch Sync_Status on Grid."),
        ("APN Oppt Status Table", "View only. Check AWS_Review_Status and deadlines. Do not add rows."),
        ("Docs", "Process guide in the APN folder if you forget a click."),
        ("APN AWS Invitations", "Sales admin (Vietnam and Indonesia). Referral accept / reject. Extra for AMs."),
    ]
    for i, (title, body) in enumerate(tables):
        col, row = i % 2, i // 2
        l = Inches(0.5 + col * 6.4)
        t = Inches(1.35 + row * 2.7)
        card(s, l, t, Inches(6.15), Inches(2.5))
        add_text(s, l + Inches(0.25), t + Inches(0.35), Inches(5.65), Inches(0.55), title, size=16, bold=True, color=NAVY)
        add_text(s, l + Inches(0.25), t + Inches(1.05), Inches(5.65), Inches(1.2), body, size=13, color=MUTE)
    footer(s)

    # 4 Jobs
    s = prs.slides.add_slide(blank)
    cream_bg(s)
    kicker(s, Inches(0.5), Inches(0.32), Inches(12), "Which job?")
    add_text(s, Inches(0.5), Inches(0.62), Inches(12), Inches(0.45), "Pick one path. Do not mix Form, Add Record send, Retry, and Update.", size=20, bold=True, color=NAVY)
    jobs = [
        (WARN_BG, "Draft", "Every AM", "Grid → + Add Record\nLeave Submit unticked"),
        (OK_BG, "Send", "Every AM", "Form → Submit\nor tick Submit on a draft"),
        (PAPER, "Change", "Every AM", "Only if Status is Approved\nGrid → Update"),
        (NAVY, "Referral", "Sales admin", "Invitations table\nNo Form"),
    ]
    for i, (bg, title, who, click) in enumerate(jobs):
        l = Inches(0.4 + i * 3.22)
        card(s, l, Inches(1.3), Inches(3.05), Inches(4.0), bg)
        tc = WHITE if bg == NAVY else NAVY
        bc = RGBColor(0xE8, 0xDF, 0xD2) if bg == NAVY else MUTE
        add_text(s, l + Inches(0.18), Inches(1.5), Inches(2.7), Inches(0.7), title, size=18, bold=True, color=tc)
        add_text(s, l + Inches(0.18), Inches(2.25), Inches(2.7), Inches(0.4), who, size=12, bold=True, color=ORANGE)
        add_text(s, l + Inches(0.18), Inches(2.75), Inches(2.7), Inches(2.1), click, size=14, color=bc)
    add_text(s, Inches(0.5), Inches(5.5), Inches(12), Inches(0.5), "Sales admin for referrals: Michael Nguyen (Vietnam) and Alvindo (Indonesia).", size=13, color=MUTE)
    pic(s, ASSETS / "apn-kt-three-doors.png", Inches(10.55), Inches(5.45), Inches(2.3), Inches(1.5))
    footer(s)

    # 5 Draft Add Record
    s = prs.slides.add_slide(blank)
    cream_bg(s)
    kicker(s, Inches(0.45), Inches(0.28), Inches(12), "Draft — Grid + Add Record")
    add_text(s, Inches(0.45), Inches(0.55), Inches(6.4), Inches(0.7), "Click + Add Record. Leave Submit off.", size=24, bold=True, color=NAVY)
    pic(s, ASSETS / "guide-grid-add-record.jpg", Inches(0.4), Inches(1.35), Inches(7.15), Inches(4.0))
    pic(s, ASSETS / "guide-add-record-button.jpg", Inches(0.4), Inches(5.45), Inches(7.15), Inches(1.45))
    steps = [
        "Open Submission → Grid.",
        "Click + Add Record on the toolbar. Not + New under the table list.",
        "Fill the row. Leave Action as Create.",
        "Leave Submit and Submit_to_AWS unticked. Sync_Status empty = draft. AWS has not received it.",
        "When ready: tick both. Set Pending if it does not move. Keep Action = Create.",
    ]
    y = Inches(1.35)
    for step in steps:
        card(s, Inches(7.75), y, Inches(5.1), Inches(0.95))
        add_text(s, Inches(7.95), y + Inches(0.12), Inches(4.7), Inches(0.75), step, size=12, color=INK)
        y += Inches(1.02)
    footer(s)

    # 6 Form
    s = prs.slides.add_slide(blank)
    cream_bg(s)
    kicker(s, Inches(0.45), Inches(0.28), Inches(12), "Send now — Form")
    add_text(s, Inches(0.45), Inches(0.55), Inches(12), Inches(0.55), "One-shot Create. Fill, then Submit.", size=24, bold=True, color=NAVY)
    pic(s, ASSETS / "guide-form-fill.jpg", Inches(0.4), Inches(1.25), Inches(6.9), Inches(5.55))
    form_steps = [
        ("Form", "Use Fill, not Edit. Fill required fields. Follow the format rules."),
        ("Submit", "Click Submit. Then open Grid."),
        ("Wait one minute", "Sync_Status: Pending becomes Success or Failed."),
        ("If Failed", "Side Peek → Sync_Message → fix that field → Retry. Leave Action = Create. Do not click Update."),
    ]
    y = Inches(1.25)
    for title, body in form_steps:
        bg = BAD_BG if title.startswith("If") else PAPER
        card(s, Inches(7.55), y, Inches(5.3), Inches(1.2), bg)
        add_text(s, Inches(7.75), y + Inches(0.12), Inches(4.95), Inches(0.32), title, size=14, bold=True, color=BAD if title.startswith("If") else NAVY)
        add_text(s, Inches(7.75), y + Inches(0.48), Inches(4.95), Inches(0.62), body, size=12, color=MUTE)
        y += Inches(1.32)
    footer(s)

    # 7 Update
    s = prs.slides.add_slide(blank)
    cream_bg(s)
    kicker(s, Inches(0.5), Inches(0.32), Inches(12), "Update")
    add_text(s, Inches(0.5), Inches(0.62), Inches(12), Inches(0.5), "Only if Status is Approved. Do not type in Status.", size=22, bold=True, color=NAVY)
    cards7 = [
        (OK_BG, OK, "Go", "Approved. Edit on Submission. Click Update. Leave Action as Update."),
        (WARN_BG, ORANGE, "Stop", "Pending submission. Do not Update. Wait on Status."),
        (BAD_BG, BAD, "Stop", "Create still Failed. Retry first. Do not Update."),
    ]
    for i, (bg, col, title, body) in enumerate(cards7):
        l = Inches(0.45 + i * 4.25)
        card(s, l, Inches(1.35), Inches(4.05), Inches(2.6), bg)
        add_text(s, l + Inches(0.22), Inches(1.5), Inches(3.6), Inches(0.45), title, size=22, bold=True, color=col)
        add_text(s, l + Inches(0.22), Inches(2.1), Inches(3.6), Inches(1.5), body, size=14, color=INK)
    card(s, Inches(0.45), Inches(4.2), Inches(7.6), Inches(2.4), PAPER)
    add_text(s, Inches(0.7), Inches(4.4), Inches(7.1), Inches(0.4), "If Update then Failed", size=16, bold=True, color=NAVY)
    add_text(s, Inches(0.7), Inches(4.9), Inches(7.1), Inches(1.4), "Read Sync_Message. Fix that field. Click Update again. Do not click Retry on an Update.", size=15, color=MUTE)
    pic(s, ASSETS / "apn-kt-close.png", Inches(8.3), Inches(4.2), Inches(4.55), Inches(2.4))
    footer(s)

    # 8 Referral
    s = prs.slides.add_slide(blank)
    cream_bg(s)
    kicker(s, Inches(0.5), Inches(0.32), Inches(12), "Referral — sales admin")
    add_text(s, Inches(0.5), Inches(0.62), Inches(12), Inches(0.5), "No Form. The invitation row already exists.", size=22, bold=True, color=NAVY)
    pic(s, ASSETS / "apn-kt-referral.png", Inches(0.4), Inches(1.3), Inches(5.5), Inches(5.3))
    refs = [
        "Open APN AWS Invitations.",
        "Find the row (country + engagement title).",
        "Set Decision to Accept (or Reject).",
        "Set Assignee to the account manager.",
        "Set Sync_Status to Pending. Then wait.",
        "AWS_Opportunity_ID appears. Status still takes about six hours.",
    ]
    y = Inches(1.3)
    for step in refs:
        card(s, Inches(6.15), y, Inches(6.7), Inches(0.78))
        add_text(s, Inches(6.4), y + Inches(0.18), Inches(6.25), Inches(0.5), step, size=13, color=INK)
        y += Inches(0.85)
    footer(s)

    # 9 Clocks
    s = prs.slides.add_slide(blank)
    cream_bg(s)
    kicker(s, Inches(0.5), Inches(0.28), Inches(12), "Two clocks and buttons")
    add_text(s, Inches(0.5), Inches(0.55), Inches(12), Inches(0.45), "Draft with ticks off does not start the 1-minute clock.", size=18, bold=True, color=NAVY)
    card(s, Inches(0.45), Inches(1.15), Inches(6.05), Inches(2.5), NAVY)
    add_text(s, Inches(0.7), Inches(1.35), Inches(5.5), Inches(0.35), "After Form Submit or Update", size=13, bold=True, color=ORANGE)
    add_text(s, Inches(0.7), Inches(1.75), Inches(5.5), Inches(1.0), "1 min", size=48, bold=True, color=WHITE)
    add_text(s, Inches(0.7), Inches(2.85), Inches(5.5), Inches(0.5), "Look at Submission → Sync_Status", size=14, color=RGBColor(0xE8, 0xDF, 0xD2))
    card(s, Inches(6.75), Inches(1.15), Inches(6.05), Inches(2.5), NAVY)
    add_text(s, Inches(7.0), Inches(1.35), Inches(5.5), Inches(0.35), "After Success create", size=13, bold=True, color=ORANGE)
    add_text(s, Inches(7.0), Inches(1.75), Inches(5.5), Inches(1.0), "6 hours", size=48, bold=True, color=WHITE)
    add_text(s, Inches(7.0), Inches(2.85), Inches(5.5), Inches(0.5), "Look at Status → AWS_Review_Status", size=14, color=RGBColor(0xE8, 0xDF, 0xD2))
    mix = [
        ("If this happened", "Click", "Leave Action as"),
        ("New Form submit Failed", "Retry", "Create"),
        ("Approved opportunity needs a change", "Update", "Update"),
        ("Update then Failed", "Update again", "Update"),
    ]
    y = Inches(3.85)
    for i, row in enumerate(mix):
        bg = NAVY if i == 0 else PAPER
        tc = WHITE if i == 0 else INK
        card(s, Inches(0.45), y, Inches(12.35), Inches(0.7), bg)
        add_text(s, Inches(0.65), y + Inches(0.18), Inches(6.2), Inches(0.4), row[0], size=13, bold=(i == 0), color=tc)
        add_text(s, Inches(7.0), y + Inches(0.18), Inches(2.6), Inches(0.4), row[1], size=13, bold=True, color=ORANGE if i else WHITE)
        add_text(s, Inches(9.8), y + Inches(0.18), Inches(2.7), Inches(0.4), row[2], size=13, bold=(i == 0), color=tc)
        y += Inches(0.75)
    footer(s)

    # 10 Format
    s = prs.slides.add_slide(blank)
    cream_bg(s)
    kicker(s, Inches(0.5), Inches(0.32), Inches(12), "Format — why rows Fail")
    add_text(s, Inches(0.5), Inches(0.62), Inches(12), Inches(0.5), "Read Sync_Message and fix that field.", size=22, bold=True, color=NAVY)
    rules = [
        ("Problem ≥ 20 characters", "Customer need must be a real sentence or Sync_Status Fails."),
        ("Phone E.164", "+601121555525  — plus, country code, no spaces."),
        ("Launched Stage", "You must have AWS Account ID. Missing ID = Failed."),
    ]
    for i, (title, body) in enumerate(rules):
        l = Inches(0.45 + i * 4.25)
        card(s, l, Inches(1.4), Inches(4.05), Inches(4.6))
        add_text(s, l + Inches(0.25), Inches(1.75), Inches(3.55), Inches(1.2), title, size=22, bold=True, color=NAVY)
        add_text(s, l + Inches(0.25), Inches(3.2), Inches(3.55), Inches(2.0), body, size=15, color=MUTE)
    footer(s)

    # 11 Go-live
    s = prs.slides.add_slide(blank)
    navy_bg(s)
    pic(s, ASSETS / "apn-kt-golive.png", Inches(6.5), 0, Inches(6.84), H)
    overlay = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(7.3), H)
    fill(overlay, NAVY)
    kicker(s, Inches(0.55), Inches(0.7), Inches(6.3), "Existing opportunities and go-live")
    add_text(s, Inches(0.55), Inches(1.2), Inches(6.3), Inches(1.4), "18 September 2026", size=36, bold=True, color=WHITE)
    add_para_box(
        s,
        Inches(0.55),
        Inches(2.8),
        Inches(6.2),
        Inches(3.4),
        [
            ("Open opportunities from 2025 to now will be moved into this Base. Do not recreate those deals.", False, 16, RGBColor(0xE8, 0xDF, 0xD2)),
            ("All GAP sales use this Lark Base from 18 September.", True, 16, WHITE),
            ("Until then you may practise Add Record. Do not tick Submit on a real customer.", False, 16, RGBColor(0xE8, 0xDF, 0xD2)),
        ],
    )
    footer(s, light=True)

    # 12 Help
    s = prs.slides.add_slide(blank)
    cream_bg(s)
    kicker(s, Inches(0.5), Inches(0.32), Inches(12), "If you are stuck")
    add_text(s, Inches(0.5), Inches(0.62), Inches(12), Inches(0.5), "Error, bug, or unclear Sync_Message?", size=24, bold=True, color=NAVY)
    card(s, Inches(0.45), Inches(1.35), Inches(6.15), Inches(4.6), NAVY)
    add_text(s, Inches(0.7), Inches(1.55), Inches(5.7), Inches(0.35), "APN TroubleShoot", size=13, bold=True, color=ORANGE)
    add_text(s, Inches(0.7), Inches(2.0), Inches(5.7), Inches(0.7), "Ask in this group. Always send three things:", size=16, color=WHITE)
    add_para_box(
        s,
        Inches(0.7),
        Inches(2.8),
        Inches(5.7),
        Inches(2.6),
        [
            ("Your Oppt ID", True, 18, WHITE),
            ("Your Status (Sync_Status / AWS_Review_Status)", True, 18, WHITE),
            ("Your Sync_Message", True, 18, WHITE),
        ],
    )
    card(s, Inches(6.85), Inches(1.35), Inches(5.95), Inches(2.15), WARN_BG)
    add_text(s, Inches(7.1), Inches(1.55), Inches(5.5), Inches(0.35), "APN flow", size=13, bold=True, color=ORANGE)
    add_text(s, Inches(7.1), Inches(2.0), Inches(5.5), Inches(0.5), "Ask Eylia", size=22, bold=True, color=NAVY)
    add_text(s, Inches(7.1), Inches(2.55), Inches(5.5), Inches(0.5), "+6016-4244515  ·  or Azib / APN champions", size=14, color=INK)
    pic(s, ASSETS / "apn-kt-help.png", Inches(6.85), Inches(3.7), Inches(5.95), Inches(2.25))
    add_text(s, Inches(0.5), Inches(6.15), Inches(12), Inches(0.4), "Do not patch AWS_Review_Status by hand.", size=14, bold=True, color=BAD)
    footer(s)

    prs.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    build()
